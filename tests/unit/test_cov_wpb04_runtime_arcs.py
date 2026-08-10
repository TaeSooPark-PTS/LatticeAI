"""wpb04 — never-taken branch directions in the runtime builders, the CLI
tunnel helpers, and the document tools.

The CLI tests replace the download, the process launch and the clock with
recorders, so nothing is fetched, spawned or waited on. The Windows-only chmod
skip is driven by patching ``sys.platform``, which keeps the line executable on
the ubuntu coverage leg.
"""

from __future__ import annotations

import itertools
import json
import sys
from pathlib import Path
from types import ModuleType, SimpleNamespace
from typing import Any, List

import pytest

import latticeai.tools as tools
from latticeai.cli import entrypoint
from latticeai.runtime.audit_runtime import build_audit_runtime
from latticeai.runtime.context_runtime import build_context_runtime
from latticeai.tools import ToolError
from latticeai.tools import documents as documents_module
from latticeai.tools.documents import create_docx, read_document

# ── audit_runtime ────────────────────────────────────────────────────────────


class _Logging:
    def __init__(self) -> None:
        self.warnings: List[Any] = []

    def warning(self, *args: Any) -> None:
        self.warnings.append(args)


def _audit(tmp_path: Path):
    audit_file = tmp_path / "audit.json"
    runtime = build_audit_runtime(audit_file=audit_file, logging=_Logging())
    return runtime, audit_file, audit_file.with_suffix(audit_file.suffix + ".jsonl")


def test_a_legacy_audit_file_that_is_not_a_list_is_ignored(tmp_path):
    """audit_runtime.py:33→35 — a JSON object where a list was expected, and
    41→37 — a JSONL line that parses to something that is not an event."""
    runtime, audit_file, jsonl = _audit(tmp_path)
    audit_file.write_text(json.dumps({"events": "not a list"}), encoding="utf-8")
    jsonl.write_text(
        "\n".join([json.dumps("a bare string"), "", json.dumps({"event_type": "chat_message"})]),
        encoding="utf-8",
    )

    events = runtime["get_audit_log"]()

    assert events == [{"event_type": "chat_message"}]


def test_an_event_with_no_payload_still_records_its_type_and_time(tmp_path):
    """audit_runtime.py:49→51 — ``_append`` called with no keyword payload."""
    runtime, _audit_file, jsonl = _audit(tmp_path)

    runtime["append_audit_event"]("server_start")

    written = [json.loads(line) for line in jsonl.read_text(encoding="utf-8").splitlines() if line]
    assert len(written) == 1
    assert written[0]["event_type"] == "server_start"
    assert written[0]["timestamp"]
    assert set(written[0]) == {"event_type", "timestamp"}


# ── context_runtime ──────────────────────────────────────────────────────────


class _Graph:
    def filter_scoped_nodes(self, matches, allowed, *, include_legacy_global=False):
        return list(matches)

    def search(self, query, limit, *, allowed_workspaces=None, include_legacy_global=False):
        return {"matches": []}

    def relationship_search(self, **_kwargs):
        return {"relationships": []}

    def vector_search(self, query, *, limit=30, min_score=0.0):
        return {"matches": []}


class _Gardener:
    def __init__(self) -> None:
        self.calls: List[Any] = []

    def get_relevant_context(self, query, *, allowed_workspaces=None, **_kw):
        self.calls.append(allowed_workspaces)
        return "정원 컨텍스트"


class _Ledger:
    def recent(self, **_kwargs):
        return [{"path": "a.html"}]


def _context(*, require_auth: bool, gardener: _Gardener, ledger: Any = None):
    return build_context_runtime(
        graph_store=_Graph(),
        ingestion_pipeline=None,
        memory_service=SimpleNamespace(recall=lambda *a, **k: {"results": []}),
        gardener=gardener,
        require_auth=require_auth,
        allowed_scopes_for_user=lambda user: {"w-" + str(user)},
        artifact_ledger=ledger,
    )


def test_search_is_unscoped_when_auth_is_off_and_the_ledger_is_injected():
    """context_runtime.py:29→34 (no auth, so no scope is derived) and 48→53
    (a ledger was supplied, so no default one is constructed)."""
    ledger = _Ledger()
    runtime = _context(require_auth=False, gardener=_Gardener(), ledger=ledger)

    payload = runtime["_scoped_hybrid_search"]("릴리스", user_email="u@example.com", workspace_id="w1")

    assert payload["matches"] == []
    assert runtime["ARTIFACT_LEDGER"] is ledger
    assert runtime["CONTEXT_ASSEMBLER"] is not None


def test_notes_context_stays_unscoped_for_an_anonymous_authenticated_caller():
    """context_runtime.py:41→43 — auth is on but the request carried neither a
    workspace nor a user, so nothing can be narrowed."""
    gardener = _Gardener()
    runtime = _context(require_auth=True, gardener=gardener, ledger=_Ledger())

    text = runtime["CONTEXT_ASSEMBLER"]._notes_context("릴리스", user_email="", workspace_id=None)

    assert text == "정원 컨텍스트"
    assert gardener.calls == [None], "no scope could be derived, so none is claimed"


# ── cli/entrypoint ───────────────────────────────────────────────────────────


def test_the_windows_download_skips_the_executable_bit(tmp_path, monkeypatch):
    """entrypoint.py:131→133 — chmod is a POSIX-only step."""
    monkeypatch.setenv("HOME", str(tmp_path))
    monkeypatch.setattr(entrypoint.shutil, "which", lambda _name: None)
    monkeypatch.setattr(sys, "platform", "win32")
    downloads: List[Any] = []

    def _urlretrieve(url, dest):
        downloads.append((url, str(dest)))
        Path(dest).write_bytes(b"cloudflared")

    monkeypatch.setattr(entrypoint.urllib.request, "urlretrieve", _urlretrieve)

    path = entrypoint._ensure_cloudflared()

    assert path.endswith("cloudflared.exe")
    assert downloads[0][0].endswith("cloudflared-windows-amd64.exe")
    assert Path(path).read_bytes() == b"cloudflared"


def test_the_tunnel_waits_for_the_url_to_appear_in_the_log(tmp_path, monkeypatch):
    """entrypoint.py:179→174 — the first poll finds no URL yet, so the loop
    goes round again. The clock and the sleep are both replaced, so this waits
    on a scripted event rather than on time passing."""
    monkeypatch.setenv("HOME", str(tmp_path))
    monkeypatch.delenv("LATTICEAI_TELEGRAM_BOT_TOKEN", raising=False)
    monkeypatch.delenv("LATTICEAI_TELEGRAM_CHAT_ID", raising=False)
    monkeypatch.setattr(entrypoint, "_ensure_cloudflared", lambda: "/usr/local/bin/cloudflared")

    launched: List[Any] = []

    def _popen(cmd, stdout=None, stderr=None):
        launched.append(cmd)
        if stdout is not None:
            stdout.close()
        return SimpleNamespace(pid=4321)

    monkeypatch.setattr(entrypoint.subprocess, "Popen", _popen)

    log_path = tmp_path / ".latticeai" / "tunnel.log"
    sleeps: List[float] = []

    def _sleep(seconds):
        sleeps.append(seconds)
        if len(sleeps) == 2:  # the tunnel finally announces itself
            log_path.write_text(
                "starting\nyour url is https://calm-lattice-42.trycloudflare.com\n",
                encoding="utf-8",
            )

    ticks = itertools.count()
    monkeypatch.setattr(entrypoint.time, "sleep", _sleep)
    monkeypatch.setattr(entrypoint.time, "time", lambda: float(next(ticks)))

    url = entrypoint._start_tunnel(4825)

    assert url == "https://calm-lattice-42.trycloudflare.com"
    assert len(sleeps) == 2, "the first poll found nothing and went round again"
    assert launched[0][:2] == ["/usr/local/bin/cloudflared", "tunnel"]


# ── tools/documents ──────────────────────────────────────────────────────────


@pytest.fixture()
def workspace(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    root = tmp_path / "agent_workspace"
    root.mkdir()
    monkeypatch.setattr(tools, "AGENT_ROOT", root)
    tools.ensure_agent_root()
    return root


def test_a_docx_with_no_title_and_no_body_is_still_a_valid_document(workspace):
    """documents.py:85→87 (no heading to add) and 89→87 (an empty block adds no
    paragraph)."""
    from docx import Document

    result = create_docx("", "\n\n   \n\n", filename="empty.docx")

    written = workspace / result["path"]
    assert written.is_file() and result["bytes"] > 0
    document = Document(str(written))
    assert [p.text for p in document.paragraphs if p.text.strip()] == []


def _fake_pptx(monkeypatch, slides: List[List[Any]]) -> None:
    module = ModuleType("pptx")

    class _Presentation:
        def __init__(self, _path: str) -> None:
            self.slides = [SimpleNamespace(shapes=shapes) for shapes in slides]

    module.Presentation = _Presentation
    monkeypatch.setitem(sys.modules, "pptx", module)


def test_reading_a_pptx_skips_shapes_that_hold_no_text(tmp_path, monkeypatch):
    """documents.py:256→255 — a picture (or any frameless shape) contributes
    nothing and the shape walk continues."""
    picture = SimpleNamespace(has_text_frame=False)
    caption = SimpleNamespace(has_text_frame=True, text_frame=SimpleNamespace(text="표지"))
    _fake_pptx(monkeypatch, [[picture, caption]])
    deck = tmp_path / "deck.pptx"
    deck.write_bytes(b"not really a zip; the parser is injected")

    meta = read_document(str(deck))

    assert meta["slides"] == 1
    assert meta["content"] == "[Slide 1]\n표지"
    assert meta["chars"] == len(meta["content"])


def test_a_supported_extension_with_no_reader_yields_an_empty_extraction(tmp_path, monkeypatch):
    """documents.py:264→270 — the extension passed the allow-list but matched
    none of the parser branches, so the result is honestly empty."""
    monkeypatch.setattr(
        documents_module,
        "_SUPPORTED_READ_EXTENSIONS",
        set(documents_module._SUPPORTED_READ_EXTENSIONS) | {".rtf"},
    )
    target = tmp_path / "note.rtf"
    target.write_text("{\\rtf1 hello}", encoding="utf-8")

    meta = read_document(str(target))

    assert meta["ext"] == ".rtf"
    assert meta["chars"] == 0
    assert meta["preview"] == "" and meta["content"] == ""


def test_an_unsupported_extension_is_still_refused(tmp_path):
    target = tmp_path / "note.rtf"
    target.write_text("{\\rtf1 hello}", encoding="utf-8")

    with pytest.raises(ToolError, match="지원하지 않는 형식"):
        read_document(str(target))
